"""Создать учебные входные файлы; существующие файлы не изменять."""
from pathlib import Path

from docx import Document
from openpyxl import Workbook


def create_examples(directory: Path) -> list[str]:
    created = []
    students = directory / "students.xlsx"
    if not students.exists():
        workbook = Workbook()
        group = workbook.active
        group.title = "ИС-21"
        for row in (("Фамилия", "Имя"), ("Иванов", "Иван"), ("Петрова", "Анна"), ("Сидоров", "Пётр")):
            group.append(row)
        group = workbook.create_sheet("ИС-22")
        for row in (("Фамилия", "Имя"), ("Смирнова", "Мария"), ("Козлов", "Алексей")):
            group.append(row)
        workbook.create_sheet("Пустая группа").append(("Фамилия", "Имя"))
        for sheet in workbook:
            sheet.column_dimensions["A"].width = 26
            sheet.column_dimensions["B"].width = 24
        workbook.save(students)
        workbook.close()
        created.append(students.name)
    tickets = directory / "tickets.docx"
    if not tickets.exists():
        document = Document()
        questions = (
            ("Что такое переменная?", "Какие типы данных есть в Python?", "Для чего нужен условный оператор?"),
            ("Что такое цикл?", "Чем отличаются for и while?", "Для чего нужны break и continue?"),
            ("Что такое функция?", "Что такое параметр функции?", "Что возвращает оператор return?"),
            ("Что такое список?", "Чем список отличается от кортежа?", "Как работает словарь?"),
            ("Что такое исключение?", "Как работают try и except?", "Зачем нужен блок finally?"),
        )
        for number, items in enumerate(questions, 1):
            document.add_paragraph(f"Билет {number}", style="Heading 1")
            for index, question in enumerate(items, 1):
                document.add_paragraph(f"{index}. {question}")
        document.save(tickets)
        created.append(tickets.name)
    return created


if __name__ == "__main__":
    print(create_examples(Path(__file__).resolve().parent))
