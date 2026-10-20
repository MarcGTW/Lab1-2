"""Генератор экзаменационных билетов. Запуск: python main.py (Windows)."""

from __future__ import annotations

import os
import random
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException


JOURNAL = Path(__file__).resolve().with_name("journal.xlsx")
HEADERS = ("Last name", "First name", "Номер билета", "Дата и время")


class ExitRequested(Exception):
    """Пользователь нажал ESC или завершил поток ввода."""


class JournalError(Exception):
    """Существующий журнал нельзя безопасно изменить."""


def read_key() -> str:
    """На Windows getwch читает Unicode-клавишу без ожидания Enter."""
    import msvcrt

    key = msvcrt.getwch()
    if key in ("\x00", "\xe0"):
        msvcrt.getwch()  # Второй код стрелки, Delete или функциональной клавиши.
        return ""
    return key


def read_line(prompt: str) -> str:
    """Посимвольный ввод: ESC работает и в фамилии, и в имени."""
    print(prompt, end="", flush=True)
    characters: list[str] = []
    while True:
        key = read_key()
        if key in ("\x1b", "\x03", "\x1a"):
            print()
            raise ExitRequested
        if key in ("\r", "\n"):
            print()
            return "".join(characters).strip()
        if key == "\b":
            if characters:
                characters.pop()
                print("\b \b", end="", flush=True)
        elif key and key.isprintable():
            characters.append(key)
            print(key, end="", flush=True)


def read_name(prompt: str) -> str:
    while True:
        value = read_line(prompt)
        if value:
            return value
        print("Поле не может быть пустым. Повторите ввод.")


def new_workbook(headers=HEADERS, widths=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Журнал"
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 30
    for column, width in (widths or (("A", 25), ("B", 25), ("C", 18), ("D", 24))):
        sheet.column_dimensions[column].width = width
    return workbook


def atomic_save(workbook, path: Path) -> None:
    """Сначала пишем полный xlsx рядом, затем атомарно заменяем журнал.

    Если процесс остановится во время записи временного файла, предыдущая
    версия журнала останется целой. os.replace на Windows отклоняет замену
    файла, открытого Excel без разрешения на удаление/замену.
    """
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=".journal-", suffix=".xlsx", dir=path.parent
    )
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w+b") as stream:
            workbook.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def save_record(path: Path, record: tuple | None = None) -> None:
    """Перечитываем журнал перед каждой записью, старые строки сохраняем."""
    workbook = None
    try:
        if path.exists():
            try:
                workbook = load_workbook(path)
            except (BadZipFile, InvalidFileException, ValueError, KeyError) as error:
                raise JournalError(
                    "Файл journal.xlsx повреждён или имеет неверный формат. "
                    "Восстановите его из резервной копии либо переименуйте."
                ) from error
            sheet = workbook.worksheets[0]
            actual_headers = tuple(sheet.cell(1, column).value for column in range(1, 5))
            if actual_headers != HEADERS:
                raise JournalError(
                    "Шапка journal.xlsx не соответствует формату журнала. "
                    "Ожидается: Last name, First name, Номер билета, Дата и время."
                )
            if record is None:
                return
        else:
            workbook = new_workbook()
            sheet = workbook.active

        if record is not None:
            sheet.append(record)
            row = sheet.max_row
            # Имена всегда текст: строка '=...' не становится формулой Excel.
            for column in (1, 2):
                sheet.cell(row, column).data_type = "s"
            sheet.cell(row, 4).number_format = "dd.mm.yyyy hh:mm:ss"
        atomic_save(workbook, path)
    finally:
        if workbook is not None:
            workbook.close()


def save_with_retry(path: Path, record: tuple | None = None) -> None:
    while True:
        try:
            save_record(path, record)
            return
        except PermissionError:
            print(
                "Не удалось записать journal.xlsx. Если файл открыт в Excel, "
                "закройте его. Также проверьте права записи в папку."
            )
        except JournalError as error:
            print(error)
        except OSError as error:
            print(f"Ошибка доступа к журналу: {error}")
            print("Проверьте доступность папки и свободное место на диске.")
        if record is not None:
            print("Текущая запись ещё не сохранена. Билет при повторе не изменится.")
        print("Исправьте проблему и нажмите Enter для повтора; ESC — выход.")
        wait_for_retry()


def wait_for_retry() -> None:
    while True:
        key = read_key()
        if key in ("\x1b", "\x03", "\x1a"):
            raise ExitRequested
        if key in ("\r", "\n"):
            return


def main() -> int:
    if os.name != "nt":
        print("Это приложение использует консоль Windows (msvcrt).")
        return 1
    print("Для выхода нажмите ESC.")
    print(f"Журнал: {JOURNAL}")
    try:
        save_with_retry(JOURNAL)
        while True:
            last_name = read_name("Last name: ")
            first_name = read_name("First name: ")
            ticket = random.randint(1, 20)
            record = (last_name, first_name, ticket, datetime.now())
            print(f"Билет № {ticket}")
            save_with_retry(JOURNAL, record)
            print("Запись сохранена.\n")
    except (ExitRequested, KeyboardInterrupt, EOFError):
        print("Работа завершена. Все ранее сохранённые записи остаются в журнале.")
        return 0


if __name__ == "__main__":
    sys.exit(main())
