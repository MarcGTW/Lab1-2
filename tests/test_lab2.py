import contextlib
import io
import logging
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

from create_examples import create_examples
from ticket_core import (Assignment, DataError, ResultJournal, Student, Ticket,
                         load_students, load_tickets, students_for_group)


class Lab2Tests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.journal = ResultJournal(self.root / "results.xlsx")
        self.student = Student("ИС-21", "Иванов", "Иван")
        self.tickets = {1: Ticket(1, ("А?", "Б?", "В?")), 7: Ticket(7, ("Г?", "Д?", "Е?"))}

    def write_doc(self, paragraphs):
        document = Document()
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        document.save(self.root / "tickets.docx")
        return self.root / "tickets.docx"

    def rows(self):
        workbook = load_workbook(self.journal.path)
        try:
            return list(workbook.active.values)
        finally:
            workbook.close()

    def test_parse_word_skips_bad_tickets_and_logs(self):
        path = self.write_doc(["Билет 1", "1. А?", "2. Б?", "3. В?",
                               "Билет 2", "1. Мало вопросов", "Билет ???", "1. Сирота",
                               "Билет 7", "1. Г?", "2. Д?", "3. Е?"])
        with self.assertLogs("ticket_core", logging.WARNING):
            tickets = load_tickets(path)
        self.assertEqual(set(tickets), {1, 7})
        self.assertEqual(tickets[7].questions, ("Г?", "Д?", "Е?"))

    def test_word_native_numbered_paragraphs(self):
        document = Document()
        document.add_paragraph("Билет 3")
        for question in ("А?", "Б?", "В?"):
            document.add_paragraph(question, style="List Number")
        path = self.root / "tickets.docx"
        document.save(path)
        self.assertEqual(load_tickets(path)[3].questions, ("А?", "Б?", "В?"))

    def test_all_invalid_and_missing_sources(self):
        with self.assertRaises(DataError):
            load_students(self.root / "students.xlsx")
        with self.assertRaises(DataError):
            load_tickets(self.root / "tickets.docx")
        with self.assertLogs("ticket_core", logging.WARNING), self.assertRaises(DataError):
            load_tickets(self.write_doc(["Билет 1", "1. Только один вопрос"]))

    def test_group_lookup_and_empty_group(self):
        create_examples(self.root)
        groups = load_students(self.root / "students.xlsx")
        self.assertEqual(students_for_group(groups, "ИС-21")[0], self.student)
        self.assertEqual(len(students_for_group(groups, "ИС-22")), 2)
        self.assertEqual(students_for_group(groups, "Пустая группа"), [])
        self.assertEqual(students_for_group(groups, "Неизвестная"), [])

    def test_repeat_uses_first_record_and_appends_history(self):
        self.journal.ensure_exists()
        with patch("ticket_core.random.choice", return_value=self.tickets[7]):
            first = self.journal.prepare(self.student, self.tickets)
        self.journal.append(first)
        original_rows = self.rows()
        # Даже если позже в журнал вручную добавлен другой номер, берём первый.
        self.journal.append(Assignment(self.student, self.tickets[1], datetime.now(), True))
        fresh_journal = ResultJournal(self.journal.path)
        with patch("ticket_core.random.choice", side_effect=AssertionError("Повтор не случаен")):
            repeated = fresh_journal.prepare(self.student, self.tickets)
        fresh_journal.append(repeated)
        rows = self.rows()
        self.assertEqual(repeated.ticket.number, 7)
        self.assertTrue(repeated.repeated)
        self.assertEqual(rows[:2], original_rows)
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[1][5], "нет")
        self.assertEqual(rows[-1][3], 7)
        self.assertEqual(rows[-1][5], "да")

    def test_same_name_in_other_group_is_new_student(self):
        self.journal.append(self.journal.prepare(self.student, self.tickets))
        other = Student("ИС-22", "Иванов", "Иван")
        self.assertFalse(self.journal.prepare(other, self.tickets).repeated)

    def test_previous_ticket_removed_is_error_not_reassignment(self):
        self.journal.append(Assignment(self.student, self.tickets[7], datetime.now(), False))
        with self.assertRaises(DataError):
            self.journal.prepare(self.student, {1: self.tickets[1]})
        self.assertEqual(len(self.rows()), 2)

    def test_failed_save_keeps_old_rows_then_retry_appends_once(self):
        self.journal.ensure_exists()
        assignment = self.journal.prepare(self.student, self.tickets)
        original = self.journal.path.read_bytes()
        with patch("main.os.replace", side_effect=PermissionError("Excel lock")):
            with self.assertRaises(PermissionError):
                self.journal.append(assignment)
        self.assertEqual(self.journal.path.read_bytes(), original)
        self.journal.append(assignment)
        self.assertEqual(len(self.rows()), 2)
        self.assertEqual(self.rows()[1][3], assignment.ticket.number)

    def test_examples_do_not_overwrite_sources(self):
        self.assertEqual(len(create_examples(self.root)), 2)
        original = {name: (self.root / name).read_bytes() for name in ("students.xlsx", "tickets.docx")}
        self.assertEqual(create_examples(self.root), [])
        for name, content in original.items():
            self.assertEqual((self.root / name).read_bytes(), content)


class GuiTests(unittest.TestCase):
    def setUp(self):
        import tkinter as tk
        from gui import TicketApp

        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name)
        create_examples(self.path)
        self.root = tk.Tk()
        self.root.withdraw()
        self.addCleanup(self.root.destroy)
        self.app = TicketApp(self.root, self.path)
        self.root.update()

    def choose_student(self):
        self.app.group.set("ИС-21")
        self.app.select_group()
        self.app.student.current(0)
        self.app.update_button()

    def test_selection_button_empty_group_and_result_close(self):
        self.root.deiconify()
        self.root.update()
        self.assertTrue(self.app.generate_button.instate(["disabled"]))
        self.choose_student()
        self.assertFalse(self.app.generate_button.instate(["disabled"]))
        self.app.generate()
        self.assertIsNotNone(self.app.result_window)
        workbook = load_workbook(self.path / "results.xlsx")
        first_number = workbook.active["D2"].value
        self.assertEqual(workbook.active.max_row, 2)
        workbook.close()
        self.assertTrue(self.app.result_window.bind("<Escape>"))
        window = self.app.result_window
        self.root.update()
        window.focus_force()
        self.root.update()
        window.event_generate("<Escape>")
        self.root.update()
        self.assertIsNone(self.app.result_window)
        self.assertTrue(self.root.winfo_exists())
        self.assertTrue(self.app.generate_button.instate(["disabled"]))
        self.choose_student()
        self.app.generate()
        workbook = load_workbook(self.path / "results.xlsx")
        self.assertEqual(workbook.active["D3"].value, first_number)
        self.assertEqual(workbook.active["F3"].value, "да")
        workbook.close()
        self.app.close_result()
        self.app.group.set("Пустая группа")
        self.app.select_group()
        self.assertEqual(tuple(self.app.student["values"]), ())
        self.assertTrue(self.app.generate_button.instate(["disabled"]))

    def test_gui_retries_same_assignment_after_lock(self):
        self.choose_student()
        real_append = self.app.journal.append
        attempts = []

        def locked_once(assignment):
            attempts.append(assignment)
            if len(attempts) == 1:
                raise PermissionError("Excel lock")
            real_append(assignment)

        with patch.object(self.app.journal, "append", side_effect=locked_once), \
             patch("gui.messagebox.askretrycancel", return_value=True):
            self.app.generate()
        self.assertEqual(len(attempts), 2)
        self.assertIs(attempts[0], attempts[1])
        self.app.close_result()

    def test_missing_sources_keep_gui_alive(self):
        (self.path / "students.xlsx").unlink()
        with patch("gui.messagebox.askretrycancel", return_value=False):
            self.app.reload()
        self.assertTrue(self.root.winfo_exists())
        self.assertTrue(self.app.generate_button.instate(["disabled"]))


if __name__ == "__main__":
    unittest.main()
