import contextlib
import ctypes
import io
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import main


class JournalTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "journal.xlsx"
        self.timestamp = datetime(2026, 10, 20, 12, 30, 45)

    def record(self, index):
        return (f"Фамилия {index}", f"Имя {index}", index, self.timestamp)

    def rows(self):
        workbook = load_workbook(self.path)
        try:
            return list(workbook.active.values)
        finally:
            workbook.close()

    def test_create_and_append_after_new_process(self):
        main.save_record(self.path)
        for index in range(1, 4):
            main.save_record(self.path, self.record(index))
        original = self.rows()
        self.assertEqual(original, [main.HEADERS] + [self.record(i) for i in range(1, 4)])
        script = (
            "from pathlib import Path; from datetime import datetime; "
            "import main, sys; "
            "[main.save_record(Path(sys.argv[1]), "
            "(f'Фамилия {i}', f'Имя {i}', i, datetime(2026,10,20,12,30,45))) "
            "for i in (4,5)]"
        )
        subprocess.run(
            [sys.executable, "-c", script, str(self.path)],
            cwd=Path(main.__file__).parent, check=True, capture_output=True,
        )
        rows = self.rows()
        self.assertEqual(len(rows), 6)
        self.assertEqual(rows[:4], original)
        self.assertEqual(rows[4:], [self.record(4), self.record(5)])

    def test_saved_record_survives_forced_termination(self):
        script = (
            "import main, sys, time; from pathlib import Path; "
            "from datetime import datetime; "
            "main.save_record(Path(sys.argv[1]), "
            "('Иванов','Иван',7,datetime(2026,10,20,12,30,45))); "
            "print('saved', flush=True); time.sleep(60)"
        )
        process = subprocess.Popen(
            [sys.executable, "-c", script, str(self.path)],
            cwd=Path(main.__file__).parent, stdout=subprocess.PIPE,
            stderr=subprocess.PIPE, text=True,
        )
        try:
            self.assertEqual(process.stdout.readline().strip(), "saved")
            process.kill()
            process.communicate(timeout=10)
            self.assertEqual(self.rows()[1], ("Иванов", "Иван", 7, self.timestamp))
        finally:
            if process.poll() is None:
                process.kill()
            process.communicate()

    def test_failed_replace_keeps_previous_file_and_retry_adds_once(self):
        main.save_record(self.path, self.record(1))
        original = self.path.read_bytes()
        replace = os.replace
        attempts = 0

        def fail_once(source, destination):
            nonlocal attempts
            attempts += 1
            if attempts == 1:
                self.assertEqual(self.path.read_bytes(), original)
                raise PermissionError("File is open in Excel")
            return replace(source, destination)

        with patch("main.os.replace", side_effect=fail_once), \
             patch("main.read_key", return_value="\r"), \
             contextlib.redirect_stdout(io.StringIO()) as output:
            main.save_with_retry(self.path, self.record(2))
        self.assertIn("закройте", output.getvalue())
        self.assertEqual(self.rows(), [main.HEADERS, self.record(1), self.record(2)])
        self.assertEqual(list(self.path.parent.glob(".journal-*.xlsx")), [])

    @unittest.skipUnless(os.name == "nt", "Windows file locking")
    def test_real_windows_file_lock(self):
        from ctypes import wintypes

        main.save_record(self.path, self.record(1))
        original = self.path.read_bytes()
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        create_file = kernel.CreateFileW
        create_file.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                               ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD,
                               wintypes.HANDLE]
        create_file.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        kernel.CloseHandle.restype = wintypes.BOOL
        # Allow reading, deny writes and deletion, as for an Excel lock.
        handle = create_file(str(self.path), 0x80000000, 1, None, 3, 0, None)
        self.assertNotEqual(handle, ctypes.c_void_p(-1).value)
        try:
            with patch("main.read_key", return_value="\x1b"), \
                 contextlib.redirect_stdout(io.StringIO()) as output:
                with self.assertRaises(main.ExitRequested):
                    main.save_with_retry(self.path, self.record(2))
            self.assertIn("закройте", output.getvalue())
            self.assertEqual(self.path.read_bytes(), original)
        finally:
            kernel.CloseHandle(handle)
        main.save_record(self.path, self.record(2))
        self.assertEqual(len(self.rows()), 3)

    def test_corrupt_file_is_not_overwritten(self):
        self.path.write_bytes(b"not an xlsx")
        with self.assertRaises(main.JournalError):
            main.save_record(self.path, self.record(1))
        self.assertEqual(self.path.read_bytes(), b"not an xlsx")

    def test_wrong_headers_are_not_overwritten(self):
        workbook = main.new_workbook()
        workbook.active["A1"] = "Wrong header"
        workbook.save(self.path)
        workbook.close()
        original = self.path.read_bytes()
        with self.assertRaises(main.JournalError):
            main.save_record(self.path, self.record(1))
        self.assertEqual(self.path.read_bytes(), original)

    def test_names_are_text_and_dates_are_dates(self):
        main.save_record(self.path, ("=1+1", " Иван ", 20, self.timestamp))
        workbook = load_workbook(self.path)
        try:
            self.assertEqual(workbook.active["A2"].data_type, "s")
            self.assertEqual(workbook.active["A2"].value, "=1+1")
            self.assertIsInstance(workbook.active["D2"].value, datetime)
        finally:
            workbook.close()


class ConsoleTests(unittest.TestCase):
    def test_blank_names_retry_trim_and_backspace(self):
        keys = list("\r   \r  Ивановв\b  \r")
        with patch("main.read_key", side_effect=keys), \
             contextlib.redirect_stdout(io.StringIO()) as output:
            self.assertEqual(main.read_name("Last name: "), "Иванов")
        self.assertEqual(output.getvalue().count("Поле не может быть пустым"), 2)

    def test_esc_during_each_name(self):
        for prompt in ("Last name: ", "First name: "):
            with self.subTest(prompt=prompt), \
                 patch("main.read_key", side_effect=list("Ива\x1b")), \
                 contextlib.redirect_stdout(io.StringIO()):
                with self.assertRaises(main.ExitRequested):
                    main.read_name(prompt)

    def test_application_flow_ticket_range_and_immediate_save(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "journal.xlsx"
            keys = iter("  Иванов  \r Иван \rПетров\rПётр\rСидоров\rСидор\r\x1b")
            observed_rows = []
            save = main.save_record

            def observe_save(journal_path, record=None):
                save(journal_path, record)
                workbook = load_workbook(journal_path)
                observed_rows.append(list(workbook.active.values))
                workbook.close()

            with patch("main.JOURNAL", path), \
                 patch("main.read_key", side_effect=lambda: next(keys)), \
                 patch("main.save_record", side_effect=observe_save), \
                 contextlib.redirect_stdout(io.StringIO()) as output:
                self.assertEqual(main.main(), 0)
            self.assertEqual([len(rows) for rows in observed_rows], [1, 2, 3, 4])
            for row in observed_rows[-1][1:]:
                self.assertTrue(1 <= row[2] <= 20)
            self.assertEqual(observed_rows[-1][1][:2], ("Иванов", "Иван"))
            self.assertEqual(output.getvalue().count("Запись сохранена."), 3)


if __name__ == "__main__":
    unittest.main()
