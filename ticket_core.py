"""Чтение источников и журнал лабораторной №2, независимо от интерфейса."""
from __future__ import annotations

import logging
import random
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from main import atomic_save, new_workbook

LOGGER = logging.getLogger(__name__)
RESULT_HEADERS = ("Группа", "Фамилия", "Имя", "Номер билета", "Дата и время", "Повтор (да/нет)")


class DataError(Exception):
    """Понятная пользователю ошибка источника или журнала."""


@dataclass(frozen=True)
class Student:
    group: str
    last_name: str
    first_name: str

    @property
    def label(self):
        return f"{self.last_name} {self.first_name}"


@dataclass(frozen=True)
class Ticket:
    number: int
    questions: tuple[str, ...]


@dataclass(frozen=True)
class Assignment:
    student: Student
    ticket: Ticket
    timestamp: datetime
    repeated: bool


def text(value) -> str:
    return "" if value is None else str(value).strip()


def load_students(path: Path) -> dict[str, list[Student]]:
    if not path.is_file():
        raise DataError(f"Не найден {path.name}. Поместите файл в папку приложения.")
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        groups = {}
        for sheet in workbook:
            students, seen = [], set()
            for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True):
                last, first = map(text, row)
                if not last and not first:
                    continue
                if not last or not first:
                    LOGGER.warning("%s: неполная запись студента %r", sheet.title, row)
                    continue
                student = Student(sheet.title, last, first)
                if student not in seen:
                    students.append(student)
                    seen.add(student)
            groups[sheet.title] = students
        return groups
    except (BadZipFile, InvalidFileException, ValueError, KeyError) as error:
        raise DataError(f"Не удалось прочитать {path.name}: неверный формат Excel.") from error
    finally:
        if workbook is not None:
            workbook.close()


def students_for_group(groups, group: str) -> list[Student]:
    return list(groups.get(group, []))


def is_word_numbered(paragraph) -> bool:
    # Word может хранить номер в XML, а не в paragraph.text.
    for properties in (paragraph._p.pPr, paragraph.style.element.pPr):
        if properties is not None and properties.numPr is not None:
            return True
    return False


def load_tickets(path: Path) -> dict[int, Ticket]:
    if not path.is_file():
        raise DataError(f"Не найден {path.name}. Поместите файл в папку приложения.")
    try:
        document = Document(path)
    except (PackageNotFoundError, BadZipFile, ValueError, KeyError) as error:
        raise DataError(f"Не удалось прочитать {path.name}: неверный формат Word.") from error
    tickets = {}
    number, questions, invalid = None, [], False

    def finish():
        if number is None:
            return
        if invalid or len(questions) < 3:
            LOGGER.warning("Билет %s пропущен: некорректный текст или меньше 3 вопросов", number)
        elif number in tickets:
            LOGGER.warning("Билет %s пропущен: номер повторяется", number)
        else:
            if len(questions) > 3:
                LOGGER.warning("Билет %s: использованы первые 3 вопроса", number)
            tickets[number] = Ticket(number, tuple(questions[:3]))

    for paragraph in document.paragraphs:
        line = paragraph.text.strip()
        if not line:
            continue
        header = re.fullmatch(r"Билет\s+([1-9]\d*)", line)
        if header:
            finish()
            number, questions, invalid = int(header[1]), [], False
            continue
        question = re.fullmatch(r"\d+\.\s*(\S.*)", line)
        if question and number is not None:
            questions.append(question[1])
        elif number is not None and is_word_numbered(paragraph):
            questions.append(line)
        else:
            LOGGER.warning("Нераспознанный абзац/заголовок: %s", line)
            # Граница некорректного билета не должна приклеить вопросы к предыдущему.
            if line.casefold().startswith("билет"):
                finish()
                number, questions, invalid = None, [], False
            else:
                invalid = True
    finish()
    if not tickets:
        raise DataError("В tickets.docx нет корректных билетов с тремя вопросами. Подробности: app.log.")
    return tickets


class ResultJournal:
    def __init__(self, path: Path):
        self.path = path

    def _open(self):
        if not self.path.exists():
            return new_workbook(RESULT_HEADERS, (("A", 20), ("B", 25), ("C", 25),
                                                 ("D", 18), ("E", 24), ("F", 22)))
        try:
            workbook = load_workbook(self.path)
        except (BadZipFile, InvalidFileException, ValueError, KeyError) as error:
            raise DataError("results.xlsx повреждён. Восстановите файл или переименуйте его.") from error
        headers = tuple(workbook.worksheets[0].cell(1, col).value for col in range(1, 7))
        if headers != RESULT_HEADERS:
            workbook.close()
            raise DataError("Неверная шапка results.xlsx. Существующие данные не изменены.")
        return workbook

    def ensure_exists(self):
        workbook = self._open()
        try:
            if not self.path.exists():
                atomic_save(workbook, self.path)
        finally:
            workbook.close()

    def prepare(self, student: Student, tickets: dict[int, Ticket]) -> Assignment:
        workbook = self._open()
        try:
            for row in workbook.worksheets[0].iter_rows(min_row=2, max_col=6, values_only=True):
                if tuple(map(text, row[:3])) == (student.group, student.last_name, student.first_name):
                    number = row[3]
                    if isinstance(number, bool) or not isinstance(number, (int, float)) or int(number) != number:
                        raise DataError(f"Неверный номер билета в первой записи: {student.label}.")
                    if int(number) not in tickets:
                        raise DataError(f"Ранее выданный билет № {int(number)} отсутствует среди корректных билетов. "
                                        "Восстановите его в tickets.docx: новый номер не назначается.")
                    return Assignment(student, tickets[int(number)], datetime.now(), True)
        finally:
            workbook.close()
        if not tickets:
            raise DataError("Нет доступных билетов.")
        return Assignment(student, random.choice(list(tickets.values())), datetime.now(), False)

    def append(self, assignment: Assignment):
        workbook = self._open()
        try:
            student = assignment.student
            sheet = workbook.worksheets[0]
            sheet.append((student.group, student.last_name, student.first_name,
                          assignment.ticket.number, assignment.timestamp,
                          "да" if assignment.repeated else "нет"))
            for col in (1, 2, 3, 6):
                sheet.cell(sheet.max_row, col).data_type = "s"
            sheet.cell(sheet.max_row, 5).number_format = "dd.mm.yyyy hh:mm:ss"
            atomic_save(workbook, self.path)
        finally:
            workbook.close()
